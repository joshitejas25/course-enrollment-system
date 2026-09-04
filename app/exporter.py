from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.database import get_connection


EXPORT_DIRECTORY = Path("exports")
EXPORT_FILE = EXPORT_DIRECTORY / "course_enrollment_report.xlsx"


def style_worksheet(worksheet):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill

    for column_cells in worksheet.columns:
        maximum_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                maximum_length = max(
                    maximum_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[column_letter].width = (
            maximum_length + 2
        )


def export_data_to_excel():
    EXPORT_DIRECTORY.mkdir(exist_ok=True)

    connection = get_connection()

    students = connection.execute(
        """
        SELECT id, name, email, phone, department, year
        FROM students
        ORDER BY id
        """
    ).fetchall()

    courses = connection.execute(
        """
        SELECT
            id,
            course_code,
            course_name,
            instructor,
            credits,
            capacity,
            prerequisite
        FROM courses
        ORDER BY id
        """
    ).fetchall()

    enrollments = connection.execute(
        """
        SELECT
            enrollments.id,
            students.name AS student_name,
            courses.course_code,
            courses.course_name,
            enrollments.enrollment_date,
            enrollments.status
        FROM enrollments
        JOIN students
            ON students.id = enrollments.student_id
        JOIN courses
            ON courses.id = enrollments.course_id
        ORDER BY enrollments.id
        """
    ).fetchall()

    dashboard_data = {
        "Total Students": connection.execute(
            "SELECT COUNT(*) FROM students"
        ).fetchone()[0],

        "Total Courses": connection.execute(
            "SELECT COUNT(*) FROM courses"
        ).fetchone()[0],

        "Total Enrollments": connection.execute(
            "SELECT COUNT(*) FROM enrollments"
        ).fetchone()[0],

        "Active Enrollments": connection.execute(
            """
            SELECT COUNT(*)
            FROM enrollments
            WHERE status = 'Enrolled'
            """
        ).fetchone()[0],

        "Waitlisted Students": connection.execute(
            """
            SELECT COUNT(*)
            FROM enrollments
            WHERE status = 'Waitlisted'
            """
        ).fetchone()[0],

        "Dropped Enrollments": connection.execute(
            """
            SELECT COUNT(*)
            FROM enrollments
            WHERE status = 'Dropped'
            """
        ).fetchone()[0],
    }

    connection.close()

    workbook = Workbook()

    # Students worksheet
    student_sheet = workbook.active
    student_sheet.title = "Students"

    student_sheet.append(
        [
            "ID",
            "Name",
            "Email",
            "Phone",
            "Department",
            "Year",
        ]
    )

    for student in students:
        student_sheet.append(list(student))

    # Courses worksheet
    course_sheet = workbook.create_sheet("Courses")

    course_sheet.append(
        [
            "ID",
            "Course Code",
            "Course Name",
            "Instructor",
            "Credits",
            "Capacity",
            "Prerequisite",
        ]
    )

    for course in courses:
        course_sheet.append(list(course))

    # Enrollments worksheet
    enrollment_sheet = workbook.create_sheet("Enrollments")

    enrollment_sheet.append(
        [
            "ID",
            "Student",
            "Course Code",
            "Course Name",
            "Enrollment Date",
            "Status",
        ]
    )

    for enrollment in enrollments:
        enrollment_sheet.append(list(enrollment))

    # Dashboard worksheet
    dashboard_sheet = workbook.create_sheet("Dashboard")

    dashboard_sheet.append(["Metric", "Value"])

    for metric, value in dashboard_data.items():
        dashboard_sheet.append([metric, value])

    # Format all worksheets
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        style_worksheet(worksheet)

    workbook.save(EXPORT_FILE)

    return EXPORT_FILE